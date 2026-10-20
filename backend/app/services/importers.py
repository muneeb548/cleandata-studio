"""File format detection and import for CSV, TSV, Excel and JSON.

All importers return text columns (or explicitly typed columns for Excel, whose
cells carry real types) so the same cleaning, preview, undo and export code
works for every format.
"""

from __future__ import annotations

import json
import os
import re
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, BinaryIO, Iterator

import pandas as pd

from app.core.errors import AppError
from app.services.csv_loader import STRING_DTYPE, LoadResult, clean_headers, load_delimited

FORMATS = ("csv", "tsv", "xlsx", "xls", "json")
EXTENSIONS = {".csv": "csv", ".tsv": "tsv", ".tab": "tsv", ".xlsx": "xlsx", ".xlsm": "xlsx", ".xls": "xls",
              ".json": "json", ".jsonl": "json", ".ndjson": "json"}
FORMAT_LABELS = {"csv": "CSV", "tsv": "TSV", "xlsx": "Excel (.xlsx)", "xls": "Excel 97-2003 (.xls)", "json": "JSON"}

_UNSUPPORTED = {
    ".pdf": "PDF documents", ".doc": "Word documents", ".docx": "Word documents", ".png": "images",
    ".jpg": "images", ".jpeg": "images", ".gif": "images", ".webp": "images", ".zip": "ZIP archives",
    ".txt": "plain text files", ".xml": "XML files", ".parquet": "Parquet files", ".ods": "OpenDocument spreadsheets",
}
SUPPORTED_HINT = "Supported formats: CSV, TSV, Excel (.xlsx, .xls) and JSON."


# --------------------------------------------------------------------------- names & detection

def validate_filename(filename: str | None) -> str:
    """Return a display-safe base filename, or raise for unsupported types."""
    if not filename:
        raise AppError("missing_filename", "The uploaded file has no name.", "Choose a data file and try again.")
    base = re.split(r"[\\/]", filename)[-1].strip()[:200]  # never trust client paths
    ext = Path(base).suffix.lower()
    if ext not in EXTENSIONS:
        kind = _UNSUPPORTED.get(ext)
        message = f'"{base}" is not a supported data file' + (f" ({kind} cannot be cleaned as tables)." if kind else ".")
        raise AppError("invalid_extension", message, SUPPORTED_HINT, status=415)
    return base


def format_from_name(filename: str) -> str:
    return EXTENSIONS[Path(filename).suffix.lower()]


def sniff_magic(f: BinaryIO) -> str | None:
    """Recognise a format from the first bytes; None means 'looks like text'."""
    f.seek(0)
    head = f.read(512)
    if head.startswith(b"PK\x03\x04"):
        return "zip"
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "ole"
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith((b"\x89PNG", b"\xff\xd8\xff", b"GIF8", b"RIFF")):
        return "image"
    return None


def resolve_format(f: BinaryIO, filename: str, requested: str | None) -> str:
    """Pick the format from the user's choice or the extension, then check the bytes agree."""
    fmt = requested or format_from_name(filename)
    if fmt not in FORMATS:
        raise AppError("invalid_format", f'Unknown format "{fmt}".', SUPPORTED_HINT)
    magic = sniff_magic(f)
    if magic in ("pdf", "image"):
        raise AppError("unsupported_content", f"This file is a {'PDF' if magic == 'pdf' else 'image'}, not a data table.",
                       SUPPORTED_HINT, status=415)
    if fmt == "xlsx" and magic != "zip":
        if magic == "ole":
            return "xls"
        raise AppError("corrupt_file", "This file is not a valid .xlsx workbook.",
                       "Open it in Excel and save it again as .xlsx, or export it as CSV.", status=415)
    if fmt == "xls" and magic != "ole":
        if magic == "zip":
            return "xlsx"  # renamed .xlsx files are common
        raise AppError("corrupt_file", "This file is not a valid .xls workbook.",
                       "Open it in Excel and save it again, or export it as CSV.", status=415)
    if fmt in ("csv", "tsv", "json") and magic in ("zip", "ole"):
        raise AppError("not_csv", "This file looks like an Excel workbook or other binary document, not text.",
                       "Rename it to .xlsx/.xls or choose the Excel format.", status=415)
    return fmt


# --------------------------------------------------------------------------- excel

def _excel_text(value: Any) -> Any:
    """Represent a typed Excel cell as text when its column is mixed."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d") if value.time() == time(0) else value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() and abs(value) < 2**53 else repr(value)
    text = str(value)
    return text


def _typed_column(values: list[Any]) -> pd.Series:
    """Keep Excel's own types when a whole column agrees; otherwise use text."""
    present = [v for v in values if v is not None and not (isinstance(v, str) and v == "")]
    if present and all(isinstance(v, datetime) for v in present):
        return pd.Series(pd.to_datetime([v if v is not None and v != "" else None for v in values]))
    if present and all(isinstance(v, int) and not isinstance(v, bool) for v in present):
        if all(abs(v) < 2**53 for v in present):
            return pd.Series(pd.array([v if isinstance(v, int) else None for v in values], dtype="Int64"))
    if present and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in present):
        return pd.Series(pd.array([float(v) if isinstance(v, (int, float)) else None for v in values], dtype="Float64"))
    texts = [_excel_text(v) for v in values]
    return pd.Series(pd.array([t if t not in (None, "") else None for t in texts], dtype=STRING_DTYPE))


def _frame_from_rows(rows: Iterator[tuple], warnings: list[str]) -> pd.DataFrame:
    header: list[Any] | None = None
    columns: list[list[Any]] = []
    n = 0
    for row in rows:
        if header is None:
            if row is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
                continue
            header = list(row)
            columns = [[] for _ in header]
            continue
        if row is None or all(v is None for v in row):
            # Keep empty rows in the middle (user can remove them); trailing ones are trimmed below.
            row = ()
        if len(row) > len(columns):
            for _ in range(len(row) - len(columns)):
                columns.append([None] * n)
                header.append(None)
        for i, col in enumerate(columns):
            col.append(row[i] if i < len(row) else None)
        n += 1
    if header is None:
        raise AppError("empty_file", "The sheet contains no data.", "Choose a sheet that has a header row.")
    # Excel often reports formatted-but-empty rows at the end; drop trailing empty rows only.
    while n and all(col[n - 1] is None for col in columns):
        for col in columns:
            col.pop()
        n -= 1
    # Drop trailing columns that have neither a header nor any value.
    while columns and header[-1] in (None, "") and all(v is None for v in columns[-1]):
        columns.pop()
        header.pop()
    names = clean_headers([_excel_text(h) or "" for h in header], warnings)
    data = {name: _typed_column(col) for name, col in zip(names, columns)}
    df = pd.DataFrame(data, index=pd.RangeIndex(n)) if data else pd.DataFrame(index=pd.RangeIndex(n))
    return df


# An .xlsx file is a ZIP archive. Refuse archives that expand to more than this
# (protects against "zip bombs" that are small to upload but huge in memory).
MAX_XLSX_EXPANDED_BYTES = 1024 * 1024 * 1024
MAX_XLSX_RATIO = 200


def _check_xlsx_archive(path: str) -> None:
    import os
    import zipfile

    try:
        with zipfile.ZipFile(path) as zf:
            expanded = sum(i.file_size for i in zf.infolist())
    except (zipfile.BadZipFile, OSError) as exc:
        raise AppError("corrupt_file", "This Excel file could not be opened. It may be damaged.",
                       "Open it in Excel, save it again as .xlsx, or export it as CSV.", status=415) from exc
    size = max(1, os.path.getsize(path))
    if expanded > MAX_XLSX_EXPANDED_BYTES or (expanded > 50 * 1024 * 1024 and expanded / size > MAX_XLSX_RATIO):
        raise AppError("file_too_large", "This Excel file expands to an unsafe size when opened.",
                       "Export the sheet as CSV and upload that instead.", status=413)


def _calamine_value(value: Any) -> Any:
    """Make python-calamine's cell values look like openpyxl's, so both readers give the same data."""
    if isinstance(value, str):
        return None if value == "" else value  # calamine reports empty cells as ""
    if isinstance(value, float) and value.is_integer() and abs(value) < 2**53:
        return int(value)  # Excel stores 7 as 7.0; openpyxl reports it as 7
    if isinstance(value, date) and not isinstance(value, datetime):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, timedelta):
        return str(value)
    return value


def _load_xlsx_calamine(path: str, sheet: str | None) -> LoadResult | None:
    """Fast .xlsx reader written in Rust (pip package "python-calamine").

    Returns None (and openpyxl is used) when it is not installed, the file name is not .xlsx/.xlsm,
    or calamine cannot read the file."""
    try:
        from python_calamine import CalamineWorkbook
    except ImportError:
        return None
    if Path(path).suffix.lower() not in (".xlsx", ".xlsm"):
        return None  # calamine picks the format from the file name; a renamed file goes to openpyxl
    warnings: list[str] = []
    try:
        wb = CalamineWorkbook.from_path(path)
        names = list(wb.sheet_names)
        if sheet and sheet not in names:
            raise AppError("unknown_sheet", f'Sheet "{sheet}" was not found.', f"Available sheets: {', '.join(names)}.")
        target = sheet or names[0]
        if not sheet and len(names) > 1:
            warnings.append(f'The workbook has {len(names)} sheets; loaded "{target}". Choose another sheet in the import options.')
        rows = wb.get_sheet_by_name(target).to_python(skip_empty_area=False)
    except AppError:
        raise
    except Exception:  # calamine could not read it: let openpyxl try, it reports damaged files clearly
        return None
    df = _frame_from_rows((tuple(_calamine_value(v) for v in row) for row in rows), warnings)
    del rows
    warnings.append("Formulas were imported as their last saved values.")
    return LoadResult(df=df, encoding=None, delimiter=None, format="xlsx", warnings=warnings, sheet=target, sheets=names)


def load_xlsx(path: str, sheet: str | None) -> LoadResult:
    _check_xlsx_archive(path)  # zip-bomb check before any reader opens the file
    if os.getenv("EXCEL_READER", "auto").strip().lower() != "openpyxl":
        fast = _load_xlsx_calamine(path, sheet)
        if fast is not None:
            return fast
    try:
        import openpyxl
        import openpyxl.xml
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:  # pragma: no cover - dependency is in requirements.txt
        raise AppError("missing_dependency", "Excel support is not installed on the server.",
                       "Run: pip install -r requirements.txt", status=501) from exc
    if not openpyxl.xml.DEFUSEDXML:
        # Without defusedxml, openpyxl is vulnerable to XML entity-expansion attacks.
        raise AppError("missing_dependency", "Excel import is disabled because the 'defusedxml' package is missing.",
                       "Run: pip install -r requirements.txt", status=501)
    warnings: list[str] = []
    handle = open(path, "rb")  # a file object, so a misnamed extension doesn't matter
    try:
        wb = openpyxl.load_workbook(handle, read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError) as exc:
        handle.close()
        raise AppError("corrupt_file", "This Excel file could not be opened. It may be damaged or password-protected.",
                       "Open it in Excel, save it again as .xlsx, or export it as CSV.", status=415) from exc
    except Exception as exc:  # zipfile.BadZipFile and similar
        handle.close()
        raise AppError("corrupt_file", "This Excel file could not be opened. It may be damaged.",
                       "Open it in Excel, save it again as .xlsx, or export it as CSV.", status=415) from exc
    try:
        names = list(wb.sheetnames)
        ws = _pick_sheet(wb, names, sheet, warnings)
        df = _frame_from_rows(ws.iter_rows(values_only=True), warnings)
    finally:
        wb.close()
        handle.close()
    warnings.append("Formulas were imported as their last saved values.")
    return LoadResult(df=df, encoding=None, delimiter=None, format="xlsx", warnings=warnings,
                      sheet=ws.title, sheets=names)


def load_xls(path: str, sheet: str | None) -> LoadResult:
    try:
        import xlrd
    except ImportError as exc:
        raise AppError("missing_dependency", "Support for old .xls files is not installed on the server.",
                       "Run: pip install -r requirements.txt (installs xlrd), or save the file as .xlsx.", status=501) from exc
    warnings: list[str] = []
    try:
        book = xlrd.open_workbook(path, on_demand=True)
    except Exception as exc:
        raise AppError("corrupt_file", "This .xls file could not be opened. It may be damaged or password-protected.",
                       "Open it in Excel and save it as .xlsx, or export it as CSV.", status=415) from exc
    try:
        names = book.sheet_names()
        target = sheet if sheet else names[0]
        if target not in names:
            raise AppError("unknown_sheet", f'Sheet "{target}" was not found.', f"Available sheets: {', '.join(names)}.")
        if not sheet and len(names) > 1:
            warnings.append(f'The workbook has {len(names)} sheets; loaded "{target}". Choose another sheet in the import options.')
        sh = book.sheet_by_name(target)

        def rows():
            for r in range(sh.nrows):
                out = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_EMPTY or cell.ctype == xlrd.XL_CELL_BLANK:
                        out.append(None)
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        out.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        out.append(bool(cell.value))
                    elif cell.ctype == xlrd.XL_CELL_NUMBER:
                        v = cell.value
                        out.append(int(v) if float(v).is_integer() else v)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        out.append(None)
                    else:
                        out.append(cell.value)
                yield tuple(out)

        df = _frame_from_rows(rows(), warnings)
    finally:
        book.release_resources()
    return LoadResult(df=df, encoding=None, delimiter=None, format="xls", warnings=warnings, sheet=target, sheets=names)


def _pick_sheet(wb, names: list[str], sheet: str | None, warnings: list[str]):
    if sheet:
        if sheet not in names:
            raise AppError("unknown_sheet", f'Sheet "{sheet}" was not found.', f"Available sheets: {', '.join(names)}.")
        return wb[sheet]
    if len(names) > 1:
        warnings.append(f'The workbook has {len(names)} sheets; loaded "{names[0]}". Choose another sheet in the import options.')
    return wb[names[0]]


# --------------------------------------------------------------------------- json

_WS = re.compile(r"[ \t\n\r]*")


class _Num(str):
    """A JSON number kept as its exact original text (so 2.50 stays 2.50)."""


def _decoder() -> json.JSONDecoder:
    return json.JSONDecoder(parse_float=_Num, parse_int=_Num, parse_constant=_Num)


def _dump(value: Any) -> str:
    """Serialise nested JSON without turning original numbers into strings."""
    if isinstance(value, _Num):
        return str(value)
    if isinstance(value, dict):
        return "{" + ",".join(f"{json.dumps(str(k), ensure_ascii=False)}:{_dump(v)}" for k, v in value.items()) + "}"
    if isinstance(value, list):
        return "[" + ",".join(_dump(v) for v in value) + "]"
    return json.dumps(value, ensure_ascii=False)


def _iter_json_array(text) -> Iterator[Any]:
    """Stream the elements of a top-level JSON array without loading the whole file."""
    decoder = _decoder()
    buf = text.read(1024 * 1024)
    pos = _WS.match(buf, 0).end()
    if pos >= len(buf) or buf[pos] != "[":
        raise ValueError("not an array")
    pos += 1
    eof = False
    expect_value = True
    while True:
        pos = _WS.match(buf, pos).end()
        if pos >= len(buf) - 1 and not eof:
            more = text.read(1024 * 1024)
            if not more:
                eof = True
            buf = buf[pos:] + more
            pos = 0
            continue
        if pos >= len(buf):
            raise ValueError("unexpected end of file")
        ch = buf[pos]
        if ch == "]":
            return
        if ch == ",":
            if expect_value:
                raise ValueError("unexpected comma")
            pos += 1
            expect_value = True
            continue
        try:
            value, end = decoder.raw_decode(buf, pos)
        except json.JSONDecodeError:
            if eof:
                raise
            more = text.read(1024 * 1024)
            if not more:
                eof = True
            buf = buf[pos:] + more
            pos = 0
            continue
        if end >= len(buf) and not eof:
            # A number might continue in the next chunk: read more and decode again.
            more = text.read(1024 * 1024)
            if more:
                buf = buf[pos:] + more
                pos = 0
                continue
            eof = True
        yield value
        pos = end
        expect_value = False


def _iter_json_lines(text) -> Iterator[Any]:
    decoder = _decoder()
    for number, line in enumerate(text, start=1):
        if line.strip():
            try:
                yield decoder.decode(line)
            except json.JSONDecodeError as exc:
                raise ValueError(f"line {number}: {exc.msg}") from exc


def _flatten(obj: dict, prefix: str = "") -> Iterator[tuple[str, Any]]:
    for key, value in obj.items():
        name = f"{prefix}{key}"
        if isinstance(value, dict) and value:
            yield from _flatten(value, f"{name}.")
        else:
            yield name, value


def _json_cell(value: Any, stats: dict) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, str):
        return str(value)
    if isinstance(value, (list, dict)):
        stats["nested"] += 1
        return _dump(value)
    return str(value)


def _frame_from_records(records: Iterator[Any], warnings: list[str]) -> pd.DataFrame:
    order: list[str] = []
    cols: dict[str, list] = {}
    stats = {"nested": 0, "scalars": 0, "arrays": 0}
    n = 0
    for rec in records:
        if isinstance(rec, dict):
            items = _flatten(rec)
        elif isinstance(rec, list):
            stats["arrays"] += 1
            items = ((f"column_{i + 1}", v) for i, v in enumerate(rec))
        else:
            stats["scalars"] += 1
            items = iter([("value", rec)])
        for key, value in items:
            col = cols.get(key)
            if col is None:
                col = cols[key] = [None] * n
                order.append(key)
            col.append(_json_cell(value, stats))
        n += 1
        for key in order:  # pad columns this record did not have
            if len(cols[key]) < n:
                cols[key].append(None)
    if stats["nested"]:
        warnings.append(f"{stats['nested']:,} value(s) were lists or objects and were kept as JSON text.")
    if any("." in k for k in order):
        warnings.append("Nested objects were flattened into columns such as parent.child.")
    if stats["arrays"]:
        warnings.append("Rows given as arrays were mapped to columns column_1, column_2, ….")
    names = clean_headers(order, warnings)
    data = {name: pd.array(cols[key], dtype=STRING_DTYPE) for name, key in zip(names, order)}
    return pd.DataFrame(data, index=pd.RangeIndex(n)) if data else pd.DataFrame(index=pd.RangeIndex(n))


def load_json(f: BinaryIO) -> LoadResult:
    import io

    from app.services.csv_loader import detect_encoding

    f.seek(0, io.SEEK_END)
    if f.tell() == 0:
        raise AppError("empty_file", "The file is empty.", "Check that you exported the data correctly and try again.")
    encoding, warnings = detect_encoding(f)
    f.seek(0)
    text = io.TextIOWrapper(f, encoding=encoding)
    try:
        start = text.read(4096).lstrip()
        text.detach()
        if not start:
            raise AppError("empty_file", "The file contains no data.", None)
        f.seek(0)
        text = io.TextIOWrapper(f, encoding=encoding)
        try:
            if start.startswith("["):
                df = _frame_from_records(_iter_json_array(text), warnings)
            elif start.startswith("{") and _looks_like_json_lines(start):
                df = _frame_from_records(_iter_json_lines(text), warnings)
                warnings.append("Read as JSON Lines (one object per line).")
            elif start.startswith("{"):
                # A single object: use the list it wraps (e.g. {"data": [...]}).
                obj = _decoder().decode(text.read())
                lists = [(k, v) for k, v in obj.items() if isinstance(v, list)] if isinstance(obj, dict) else []
                if len(lists) == 1:
                    key, rows = lists[0]
                    warnings.append(f'Loaded the records inside the "{key}" field.')
                    df = _frame_from_records(iter(rows), warnings)
                else:
                    df = _frame_from_records(iter([obj]), warnings)
                    warnings.append("The file contains a single JSON object, loaded as one row.")
            else:
                raise AppError("invalid_json", "This JSON file does not contain a table of records.",
                               "Use an array of objects, e.g. [{\"name\": \"Ali\", \"age\": 25}, …], or JSON Lines.")
        except (ValueError, RecursionError) as exc:
            if isinstance(exc, AppError):
                raise
            detail = str(exc).splitlines()[0][:120]
            raise AppError("invalid_json", "This file is not valid JSON.",
                           f"Check the file in a JSON validator ({detail}).") from exc
        except MemoryError as exc:
            raise AppError("out_of_memory", "There is not enough memory to load this JSON file.",
                           "Convert it to JSON Lines or CSV, or split it into smaller files.", status=507) from exc
    finally:
        try:
            text.detach()
        except ValueError:
            pass
    if len(df) == 0:
        warnings.append("The JSON file contains no records.")
    return LoadResult(df=df, encoding=encoding, delimiter=None, format="json", warnings=warnings)


def _looks_like_json_lines(start: str) -> bool:
    first_line = start.split("\n", 1)
    if len(first_line) < 2:
        return False
    try:
        json.loads(first_line[0])
    except json.JSONDecodeError:
        return False
    return first_line[1].lstrip().startswith("{")


# --------------------------------------------------------------------------- dispatcher

def load_file(path: str, filename: str, requested_format: str | None = None, sheet: str | None = None) -> LoadResult:
    with open(path, "rb") as f:
        fmt = resolve_format(f, filename, requested_format)
        if fmt in ("csv", "tsv"):
            return load_delimited(f, delimiter="\t" if fmt == "tsv" else None, fmt=fmt)
        if fmt == "json":
            return load_json(f)
    if fmt == "xlsx":
        return load_xlsx(path, sheet)
    return load_xls(path, sheet)
