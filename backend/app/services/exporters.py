"""Write the working dataset to CSV or Excel files on disk (streamed, not built in memory)."""

from __future__ import annotations

import re

import numpy as np
import pandas as pd

from app.core.errors import AppError

EXCEL_MAX_ROWS = 1_048_576  # including the header row
EXCEL_MAX_COLS = 16_384
EXCEL_MAX_CELL_CHARS = 32_767
CHUNK_ROWS = 50_000
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def escape_formulas(df: pd.DataFrame) -> pd.DataFrame:
    """Prefix text that a spreadsheet would treat as a formula with an apostrophe
    (OWASP CSV-injection guidance). Numbers such as -3.5 are left unchanged."""
    from app.services.profiling import numeric_view

    out = df.copy(deep=False)
    for col in df.columns:
        s = df[col]
        if not pd.api.types.is_string_dtype(s):
            continue
        risky = s.str.startswith(FORMULA_PREFIXES).fillna(False).astype(bool) & numeric_view(s).isna()
        if risky.any():
            out[col] = s.where(~risky, "'" + s)
    return out


# Highlight codes for Excel downloads (0 = no colour).
EMPTY, PROBLEM, DUPLICATE = 1, 2, 3
HIGHLIGHT_COLOURS = {EMPTY: "E7E6E6", PROBLEM: "FFE699", DUPLICATE: "F8CBAD"}
HIGHLIGHT_KEY = [
    ("Grey", EMPTY, "Empty cell."),
    ("Yellow", PROBLEM, 'Value that needs a look: "N/A"-style text, only spaces, extra spaces at the start or end, '
                        "an invalid email address, or a row marked in a review column."),
    ("Orange", DUPLICATE, "Row that is an exact copy of an earlier row."),
]
_EMAIL_NAME = re.compile(r"e-?mail", re.IGNORECASE)
_REVIEW_NAME = re.compile(r"^(needs_review|duplicate_group)$")


def highlight_codes(df: pd.DataFrame) -> np.ndarray:
    """One code per cell (rows x columns) saying which colour the cell gets in an Excel download."""
    from app.services.operations import EMAIL_RE
    from app.services.profiling import PLACEHOLDER_TOKENS

    rows, cols = df.shape
    codes = np.zeros((rows, cols), dtype=np.int8)
    if not rows or not cols:
        return codes
    for j, name in enumerate(df.columns):
        s = df.iloc[:, j]
        codes[s.isna().to_numpy(), j] = EMPTY
        if not pd.api.types.is_string_dtype(s) or pd.api.types.is_bool_dtype(s):
            continue
        stripped = s.str.strip()
        bad = (stripped.str.lower().isin(PLACEHOLDER_TOKENS) | (stripped == "") | (s != stripped)).fillna(False)
        if _EMAIL_NAME.search(str(name)):
            bad |= (s.notna() & ~stripped.str.fullmatch(EMAIL_RE).fillna(False)).fillna(False)
        if _REVIEW_NAME.match(str(name)):
            bad = s.notna()  # the cell that explains why the row needs attention
        codes[bad.to_numpy(dtype=bool), j] = PROBLEM
    duplicate = df.duplicated(keep="first").to_numpy()
    codes[duplicate, :] = DUPLICATE
    return codes


def write_export(df: pd.DataFrame, path: str, fmt: str, *, bom: bool = False, protect_formulas: bool = False,
                 report_lines: list[str] | None = None, highlight: bool = False) -> None:
    if fmt == "csv":
        write_csv(escape_formulas(df) if protect_formulas else df, path, bom=bom)
    elif fmt == "xlsx":
        write_xlsx(df, path, report_lines=report_lines, highlight=highlight)
    else:  # pragma: no cover - validated by the caller
        raise AppError("invalid_format", f'Cannot export as "{fmt}".', None)


def write_csv(df: pd.DataFrame, path: str, *, bom: bool = False) -> None:
    encoding = "utf-8-sig" if bom else "utf-8"
    if df.shape[1] == 0:
        # A dataset with no columns exports as an empty file instead of a stray "".
        with open(path, "w", encoding=encoding, newline=""):
            pass
        return
    # pandas writes in chunks, so the whole file never exists as one string.
    df.to_csv(path, index=False, lineterminator="\n", encoding=encoding, chunksize=CHUNK_ROWS)


def _python_values(col: pd.Series) -> list:
    """Column values as plain Python objects, with missing values as None."""
    if pd.api.types.is_datetime64_any_dtype(col):
        values = [None if pd.isna(v) else v.to_pydatetime() for v in col]
        return values
    out = col.astype(object)
    return out.where(col.notna(), None).tolist()


def write_xlsx(df: pd.DataFrame, path: str, *, report_lines: list[str] | None = None, highlight: bool = False) -> None:
    """The data on a "Cleaned data" sheet, plus an optional "Cleaning report" sheet for the client.

    highlight=True colours cells that still need a look (see HIGHLIGHT_KEY) and adds a "Colour key" sheet.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.cell.cell import Cell
        from openpyxl.styles import PatternFill
    except ImportError as exc:  # pragma: no cover - dependency is in requirements.txt
        raise AppError("missing_dependency", "Excel export is not installed on the server.",
                       "Run: pip install -r requirements.txt", status=501) from exc

    rows, cols = df.shape
    if rows + 1 > EXCEL_MAX_ROWS:
        raise AppError("excel_too_many_rows",
                       f"Excel sheets hold at most {EXCEL_MAX_ROWS - 1:,} data rows; this dataset has {rows:,}.",
                       "Download as CSV instead, or remove rows first.", status=422)
    if cols > EXCEL_MAX_COLS:
        raise AppError("excel_too_many_columns",
                       f"Excel sheets hold at most {EXCEL_MAX_COLS:,} columns; this dataset has {cols:,}.",
                       "Download as CSV instead.", status=422)

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Cleaned data")

    fills = {code: PatternFill("solid", start_color=colour, end_color=colour) for code, colour in HIGHLIGHT_COLOURS.items()}
    codes = highlight_codes(df) if highlight else None

    def coloured(value, code: int, sheet=None):
        cell = value if isinstance(value, Cell) else WriteOnlyCell(sheet or ws, value=value)
        cell.fill = fills[code]
        return cell

    def safe(value, column: str):
        if isinstance(value, str):
            if len(value) > EXCEL_MAX_CELL_CHARS:
                raise AppError("excel_cell_too_long",
                               f'A value in "{column}" is longer than Excel\'s {EXCEL_MAX_CELL_CHARS:,}-character cell limit.',
                               "Download as CSV instead, so nothing is cut off.", status=422)
            if _ILLEGAL_XML.search(value):
                raise AppError("excel_illegal_characters",
                               f'Column "{column}" contains control characters that Excel files cannot store.',
                               "Remove them with Find and replace, or download as CSV.", status=422)
            if value.startswith("="):
                # Store as text so Excel never runs it as a formula.
                cell = WriteOnlyCell(ws, value=value)
                cell.data_type = "s"
                return cell
        return value

    names = [str(c) for c in df.columns]
    try:
        ws.append([safe(n, n) for n in names])
        for start in range(0, rows, CHUNK_ROWS):
            chunk = df.iloc[start:start + CHUNK_ROWS]
            columns = [_python_values(chunk.iloc[:, i]) for i in range(cols)]
            if codes is None:
                for row in zip(*columns) if cols else ():
                    ws.append([safe(v, names[i]) for i, v in enumerate(row)])
            else:
                for r, row in enumerate(zip(*columns) if cols else ()):
                    row_codes = codes[start + r]
                    ws.append([coloured(safe(v, names[i]), row_codes[i]) if row_codes[i] else safe(v, names[i])
                               for i, v in enumerate(row)])
    except BaseException:
        try:
            ws.close()  # finish the half-written sheet stream cleanly before giving up
        except Exception:
            pass
        raise
    if codes is not None:
        key = wb.create_sheet("Colour key")
        key.append(["Colour", "Meaning", "Cells"])
        for label, code, meaning in HIGHLIGHT_KEY:
            key.append([coloured(label, code, key), meaning, int((codes == code).sum())])
    if report_lines:
        report = wb.create_sheet("Cleaning report")
        for line in report_lines:
            cell = WriteOnlyCell(report, value=_ILLEGAL_XML.sub("", line)[:EXCEL_MAX_CELL_CHARS])
            cell.data_type = "s"  # always text, never a formula
            report.append([cell])
    wb.save(path)
