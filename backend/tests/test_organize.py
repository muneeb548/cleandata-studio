"""Organize columns, sort rows, range checks, numbers written in words and coloured Excel downloads."""

import openpyxl
import pandas as pd
import pytest

from app.core.errors import AppError
from app.schemas.operations import CleanNumbers, FlagRows, OrganizeColumns, SortRows
from app.services import exporters
from app.services.operations import is_destructive, run_operation, words_to_number

CSV = (
    "id,name,age,joined,score\n"
    "3,zara,25,2024-03-01,10\n"
    "1,Ali,120,2023-01-05,\n"
    "2,bilal,-5,2025-02-01,7\n"
    "4,Aisha,thirty,2024-01-01,7\n"
    "5,,,,1\n"
)


@pytest.fixture
def df(frame):
    return frame(CSV)


# ----------------------------------------------------------------- organize columns

def test_rename_drop_and_reorder(df):
    op = OrganizeColumns(type="organize_columns", rename=[{"column": "name", "to": "Full Name"}], drop=["score"],
                         order=["name", "id"])
    res = run_operation(df, op)
    assert list(res.df.columns) == ["Full Name", "id", "age", "joined"]
    assert res.columns_removed == 1 and res.changed
    assert res.df["Full Name"].tolist()[:2] == ["zara", "Ali"]
    assert res.summary == "Removed 1 column, renamed 1 column and changed the column order."
    assert is_destructive(op) and not is_destructive(OrganizeColumns(type="organize_columns", order=["age"]))


@pytest.mark.parametrize("params,code", [
    ({"rename": [{"column": "name", "to": "id"}]}, "name_taken"),
    ({"drop": ["id", "name", "age", "joined", "score"]}, "no_columns_left"),
    ({"rename": [{"column": "nope", "to": "x"}]}, "unknown_columns"),
    ({"drop": ["name"], "order": ["name"]}, "conflicting_columns"),
    ({"drop": ["name"], "rename": [{"column": "name", "to": "x"}]}, "conflicting_columns"),
    ({"rename": [{"column": "name", "to": "a"}, {"column": "name", "to": "b"}]}, "conflicting_rename"),
])
def test_organize_refuses_conflicts(df, params, code):
    with pytest.raises(AppError) as err:
        run_operation(df, OrganizeColumns(type="organize_columns", **params))
    assert err.value.code == code


def test_swapping_two_names_is_allowed(df):
    op = OrganizeColumns(type="organize_columns", rename=[{"column": "id", "to": "name"}, {"column": "name", "to": "id"}])
    res = run_operation(df, op)
    assert list(res.df.columns)[:2] == ["name", "id"] and res.df["name"].tolist()[0] == "3"


def test_nothing_to_organize(df):
    res = run_operation(df, OrganizeColumns(type="organize_columns", order=["id"]))
    assert not res.changed and "Nothing to change" in res.summary


# ----------------------------------------------------------------- sort rows

def test_sort_text_ignores_capitals_and_keeps_row_numbers(df):
    res = run_operation(df, SortRows(type="sort_rows", by=[{"column": "name"}]))
    assert res.df["name"].tolist()[:4] == ["Aisha", "Ali", "bilal", "zara"] and pd.isna(res.df["name"].iloc[4])
    assert list(res.df.index) == [3, 1, 2, 0, 4]  # original row numbers stay with their rows
    assert res.changed and "A to Z" in res.summary


def test_sort_numbers_as_numbers_with_tie_breaker(df):
    res = run_operation(df, SortRows(type="sort_rows", by=[{"column": "score", "direction": "desc"}, {"column": "id"}]))
    assert res.df["id"].tolist() == ["3", "2", "4", "5", "1"]  # 10, 7 (id 2), 7 (id 4), 1, empty
    assert "largest first" in res.summary


def test_sort_dates_with_missing_first(df):
    res = run_operation(df, SortRows(type="sort_rows", by=[{"column": "joined"}], missing="first"))
    values = res.df["joined"].tolist()
    assert pd.isna(values[0]) and values[1:] == ["2023-01-05", "2024-01-01", "2024-03-01", "2025-02-01"]


def test_already_sorted_is_not_a_change(df):
    res = run_operation(df, SortRows(type="sort_rows", by=[{"column": "joined"}]))
    res2 = run_operation(res.df, SortRows(type="sort_rows", by=[{"column": "joined"}]))
    assert res.changed and not res2.changed


def test_sorted_data_is_exported_in_that_order(fresh_store):
    from app.services import workspace

    sid = workspace.create_session(fresh_store, "s.csv", CSV.encode())["session_id"]
    session = fresh_store.get(sid)
    workspace.apply_operation(session, SortRows(type="sort_rows", by=[{"column": "id"}]))
    data, _ = workspace.export_csv(session)
    assert [line.split(",")[0] for line in data.decode().splitlines()[1:]] == ["1", "2", "3", "4", "5"]


# ----------------------------------------------------------------- range checks

def test_out_of_range_numbers_and_dates(df):
    res = run_operation(df, FlagRows(type="flag_rows", rules=[
        {"column": "age", "check": "out_of_range", "min": "0", "max": "100"},
        {"column": "joined", "check": "out_of_range", "min": "2024-01-01"},
    ]))
    assert res.df["needs_review"].fillna("").tolist()[:3] == ["", "age out of range; joined out of range", "age out of range"]
    assert pd.isna(res.df["needs_review"].iloc[3])  # "thirty" is not a number: that is the not_number check's job


@pytest.mark.parametrize("rule", [
    {"column": "age", "check": "out_of_range"},
    {"column": "age", "check": "out_of_range", "min": "abc"},
    {"column": "age", "check": "out_of_range", "min": "0", "max": "2024-01-01"},
])
def test_invalid_ranges(df, rule):
    with pytest.raises(AppError) as err:
        run_operation(df, FlagRows(type="flag_rows", rules=[rule]))
    assert err.value.code == "invalid_range"


# ----------------------------------------------------------------- numbers in words

@pytest.mark.parametrize("text,number", [
    ("thirty", 30), ("Twenty-Five", 25), ("two hundred and five", 205), ("1.5 million", 1_500_000),
    ("minus ten", -10), ("one thousand two hundred", 1200), ("zero", 0), ("thirty one", 31),
    ("hello", None), ("and", None), ("", None), ("thirty apples", None),
])
def test_words_to_number(text, number):
    assert words_to_number(text) == number


def test_clean_numbers_reads_words_and_says_so(df):
    res = run_operation(df, CleanNumbers(type="clean_numbers", columns=["age"]))
    assert res.df["age"].tolist()[:4] == [25.0, 120.0, -5.0, 30.0]
    assert any('"thirty"' in w for w in res.warnings)


# ----------------------------------------------------------------- coloured Excel

def test_highlight_codes(frame):
    data = frame("name,email,city\nAli,ali@x.com,Lahore\n Sara ,bad-email,N/A\nAli,ali@x.com,Lahore\nZoe,,x\n")
    codes = exporters.highlight_codes(data).tolist()
    assert codes == [[0, 0, 0], [2, 2, 2], [3, 3, 3], [0, 1, 0]]


def test_coloured_excel_download(frame, tmp_path):
    data = frame("name,email\nAli,ali@x.com\n Sara ,bad\nAli,ali@x.com\n")
    path = tmp_path / "out.xlsx"
    exporters.write_xlsx(data, str(path), highlight=True)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Cleaned data", "Colour key"]
    ws = wb["Cleaned data"]
    assert ws["A2"].fill.fill_type is None
    assert ws["B3"].fill.fgColor.rgb.endswith(exporters.HIGHLIGHT_COLOURS[exporters.PROBLEM])
    assert ws["A4"].fill.fgColor.rgb.endswith(exporters.HIGHLIGHT_COLOURS[exporters.DUPLICATE])
    key = [[c.value for c in row] for row in wb["Colour key"].iter_rows(min_row=2)]
    assert [k[0] for k in key] == ["Grey", "Yellow", "Orange"] and [k[2] for k in key] == [0, 2, 2]


def test_plain_excel_has_no_colour_key(frame, tmp_path):
    path = tmp_path / "plain.xlsx"
    exporters.write_xlsx(frame("a\n1\n"), str(path))
    assert openpyxl.load_workbook(path).sheetnames == ["Cleaned data"]


def test_excel_download_is_readable(frame, tmp_path):
    """Wide enough columns, real numbers for IDs, dates without 00:00:00, and codes with a leading zero kept as text."""
    import datetime as dt

    data = frame("order_id,customer,phone,zip,price,joined\n"
                 "1001,Emma Watson-Smithington,0300-1234567,04101,12.50,2026-01-15\n"
                 "1002,Ali,+92 321 7654321,10001,7,2026-01-16\n")
    data["joined"] = pd.to_datetime(data["joined"])
    path = tmp_path / "readable.xlsx"
    exporters.write_xlsx(data, str(path))
    ws = openpyxl.load_workbook(path)["Cleaned data"]
    assert [c.value for c in ws[2]] == [1001, "Emma Watson-Smithington", "0300-1234567", "04101", 12.5,
                                        dt.datetime(2026, 1, 15)]
    assert ws["F2"].number_format == "yyyy-mm-dd"
    assert ws.column_dimensions["B"].width >= len("Emma Watson-Smithington")
    assert ws.column_dimensions["A"].width >= 8 and ws.freeze_panes == "A2"
