"""Near-duplicates, Quick Clean and the change report."""

import pandas as pd
import pytest

from app.core.errors import AppError
from app.schemas.operations import NearDuplicates, Operation
from app.services import operations as ops
from app.services import workspace
from pydantic import TypeAdapter

OP = TypeAdapter(Operation)


def values(series):
    return [None if pd.isna(v) else v for v in series.tolist()]


# ----------------------------------------------------------------- near duplicates

PEOPLE = (
    "name,email,city\n"
    "Ali Khan,ali@gmail.com,Lahore\n"
    "ali  khan.,ALI@gmail.com ,lahore\n"
    "Sara Ahmed,sara@gmail.com,Karachi\n"
    "Jon Smith,jon.smith@gmail.com,London\n"
    "John Smith,john.smith@gmail.com,London\n"
    ",,Quetta\n"
)


def test_normalized_near_duplicates_are_flagged(frame):
    df = frame(PEOPLE)
    res = ops.run_operation(df, NearDuplicates(type="near_duplicates", columns=["name", "email"]))
    assert values(res.df["duplicate_group"]) == ["Group 1", "Group 1", None, None, None, None]
    assert res.rows_removed == 0 and not ops.is_destructive(NearDuplicates(type="near_duplicates", columns=["name"]))


def test_similar_spelling_finds_more(frame):
    res = ops.run_operation(frame(PEOPLE), NearDuplicates(type="near_duplicates", columns=["name", "email"], match="similar"))
    assert values(res.df["duplicate_group"])[3:5] == ["Group 2", "Group 2"]


def test_remove_keeps_first_of_each_group(frame):
    op = NearDuplicates(type="near_duplicates", columns=["name"], action="remove")
    res = ops.run_operation(frame(PEOPLE), op)
    assert res.rows_removed == 1 and 1 not in res.df.index and 0 in res.df.index
    assert ops.is_destructive(op)


def test_empty_keys_are_never_duplicates(frame):
    res = ops.run_operation(frame("name\n\n\nx\n"), NearDuplicates(type="near_duplicates", columns=["name"]))
    assert not res.changed


def test_group_column_cannot_be_compared_column(frame):
    with pytest.raises(AppError):
        ops.run_operation(frame(PEOPLE), NearDuplicates(type="near_duplicates", columns=["name"], column_name="name"))


# ----------------------------------------------------------------- quick clean

SALES = (
    "order_id,customer,city,status,email,phone,country,order_date\n"
    "1,Ali Khan,Lahore,Paid,ali@gmail.com,0300-1234567,Pakistan,15/01/2026\n"
    "2,Sara,lahore ,paid,SARA@gmail.com,+44 20 7031 3000,UK,2026-01-16\n"
    "3,Zoya,LAHORE,PAID,zoya@gmail.com,(650) 253-0000,USA,Jan 17 2026\n"
    "4,Bilal,Karachi,N/A,bilal@gmail.com,,U.S.A.,18/01/2026\n"
    "4,Bilal,Karachi,N/A,bilal@gmail.com,,U.S.A.,18/01/2026\n"
    "5,Omar,Karachi,Pending,omar@gmail.com,0300-7654321,Pakistan,\n"
)


@pytest.fixture
def session(frame, fresh_store):
    s, _ = fresh_store.create(filename="sales.csv", file_size=1, encoding="utf-8", delimiter=",",
                              import_warnings=[], original=frame(SALES))
    return s


def test_plan_proposes_safe_fixes_only(session):
    plan = workspace.quick_clean_plan(session)
    kinds = [s["operation"]["type"] for s in plan["steps"]]
    assert kinds[0] == "remove_duplicates" and plan["steps"][0]["selected"]  # order_id makes copies safe to remove
    for expected in ("trim_whitespace", "map_values", "clean_email", "standardize_country", "standardize_phone", "convert_type"):
        assert expected in kinds
    assert "missing_values" not in kinds  # never fills or deletes missing data on its own
    assert any("missing value" in d for d in plan["needs_decision"])


def test_preview_and_apply_selected_steps(session):
    plan = workspace.quick_clean_plan(session)
    chosen = [OP.validate_python(s["operation"]) for s in plan["steps"] if s["selected"]]
    preview = workspace.preview_quick_clean(session, chosen)
    assert preview["blocked"] is None and preview["rows_after"] == 5
    assert len(session.history) == 0  # preview stores nothing
    with pytest.raises(AppError):
        workspace.apply_quick_clean(session, chosen, "wrong-token")
    result = workspace.apply_quick_clean(session, chosen, preview["preview_token"])
    assert result["applied"] == len(session.history) >= 5
    df = session.current
    assert set(df["city"].dropna()) == {"Lahore", "Karachi"}
    assert set(df["status"].dropna()) == {"Paid", "Pending"}  # "N/A" became empty
    assert values(df["country"]) == ["Pakistan", "United Kingdom", "United States", "United States", "Pakistan"]
    assert values(df["phone"])[:3] == ["+92 300 1234567", "+44 20 7031 3000", "+1 650-253-0000"]
    assert pd.api.types.is_datetime64_any_dtype(df["order_date"])


def test_change_report_message(session):
    plan = workspace.quick_clean_plan(session)
    chosen = [OP.validate_python(s["operation"]) for s in plan["steps"] if s["selected"]]
    workspace.apply_quick_clean(session, chosen, workspace.preview_quick_clean(session, chosen)["preview_token"])
    report = workspace.change_report(session)
    assert report["rows_before"] == 6 and report["rows_after"] == 5
    assert "Removed 1 duplicate row" in report["message"]
    assert len(report["steps"]) == len(session.history)
    assert report["message"].startswith("Hi!")


def test_excel_export_can_include_the_report_sheet(session):
    import os

    from openpyxl import load_workbook

    plan = workspace.quick_clean_plan(session)
    chosen = [OP.validate_python(s["operation"]) for s in plan["steps"] if s["selected"]]
    workspace.apply_quick_clean(session, chosen, workspace.preview_quick_clean(session, chosen)["preview_token"])
    path, name, _ = workspace.export_file(session, "out", "xlsx", include_report=True)
    try:
        wb = load_workbook(path, read_only=True)
        assert wb.sheetnames == ["Cleaned data", "Cleaning report"]
        lines = [row[0] for row in wb["Cleaning report"].iter_rows(values_only=True)]
        assert lines[0].startswith("Hi!") and any("duplicate row" in (line or "") for line in lines)
        wb.close()
    finally:
        os.remove(path)
    plain, _, _ = workspace.export_file(session, "out", "xlsx")
    try:
        wb = load_workbook(plain, read_only=True)
        assert wb.sheetnames == ["Cleaned data"]  # the report sheet is only added when asked for
        wb.close()
    finally:
        os.remove(plain)


def test_capitalised_spelling_wins_even_when_less_common():
    from app.services.quick_clean import _spelling_mappings

    s = pd.Series(["lahore", "lahore", "Lahore", "LAHORE", "USA", "usa", "usa", "paid", "paid", "Paid"], dtype="string")
    mapping = {m["value"]: m["to"] for m in _spelling_mappings(s)}
    assert mapping["lahore"] == "Lahore"
    assert mapping["usa"] == "USA"  # short acronym in capitals is a proper spelling
    assert mapping["paid"] == "Paid"


def test_quick_clean_removes_completely_empty_rows(frame):
    from app.services import quick_clean

    df = frame("id,name,email\n1,Ali,a@x.com\n,,\n2,Sara,s@x.com\n")
    plan = quick_clean.plan(df)
    first = plan["steps"][0]
    assert first["operation"]["type"] == "remove_empty" and first["selected"]
    assert not any("missing" in d for d in plan["needs_decision"])  # the empty row no longer counts as missing data


def test_quick_clean_converts_amounts_and_trims_header_spaces(frame):
    from app.services import quick_clean, workspace  # noqa: F401
    from app.services.operations import run_operation

    df = frame("Name ,Salary\nAli,50000\nSara,62000\nJohn,\"$1,250.50\"\nMia,700\n")
    plan = quick_clean.plan(df)
    titles = [s["title"] for s in plan["steps"]]
    assert 'Turn the amounts in "Salary" into numbers' in titles
    assert titles[-1] == "Remove spaces around column names"
    work = df
    for s in plan["steps"]:
        work = run_operation(work, OP.validate_python(s["operation"])).df
    assert list(work.columns) == ["Name", "Salary"]
    assert work["Salary"].tolist() == [50000.0, 62000.0, 1250.5, 700.0]


def test_mixed_currencies_are_not_ticked(frame):
    from app.services import quick_clean

    df = frame('price\n"$1,250"\n"€300"\n"$20"\n"Rs 500"\n')
    steps = [s for s in quick_clean.plan(df)["steps"] if s["operation"]["type"] == "clean_numbers"]
    assert steps and steps[0]["selected"] is False and "mixes currencies" in steps[0]["note"]


def test_capitals_are_fixed_even_in_a_mostly_unique_column(frame):
    """A city list where most values appear once still gets "lahore" → "Lahore" (the demo file case)."""
    from app.services.quick_clean import plan

    df = frame("order_id,city\n1,Lahore\n2,London\n3,Berlin\n4,lahore\n5,Dublin\n6,Karachi\n7,london\n")
    steps = {s["title"]: s for s in plan(df)["steps"]}
    step = steps['Use one spelling for each value in "city"']
    assert step["selected"]
    assert {m["value"]: m["to"] for m in step["operation"]["mappings"]} == {"lahore": "Lahore", "london": "London"}
    assert not any("order_id" in t for t in steps)  # ID-like columns are never touched


def test_demo_file_has_no_capitalisation_problems_left():
    from pathlib import Path

    from app.services import profiling
    from app.services.csv_loader import load_csv
    from app.services.operations import run_operation
    from app.services.quick_clean import _OPERATION, plan

    df = load_csv((Path(__file__).resolve().parents[2] / "demo_customers.csv").read_bytes()).df
    for step in plan(df)["steps"]:
        if step["selected"]:
            df = run_operation(df, _OPERATION.validate_python(step["operation"])).df
    left = [i["id"] for i in profiling.detect_issues(df) if i["severity"] != "info"]
    assert not [i for i in left if i.startswith(("case:", "dates:", "placeholder:", "whitespace:", "emails:"))], left
